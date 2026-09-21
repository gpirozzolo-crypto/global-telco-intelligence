from __future__ import annotations

import io
import re
from datetime import date

from openpyxl import load_workbook

from .base import PipelineContext, finish_run, one, start_run, utcnow
from .arcep_load import _safe_float, _upsert_obs, _upsert_raw

DATASET_API = "https://www.data.gouv.fr/api/1/datasets/le-marche-du-haut-et-tres-haut-debit-fixe-deploiements/"


def _quarter(title: str):
    m = re.search(r"(20\d{2})\s*[TQ]([1-4])", title or "", re.I)
    if not m:
        return None
    y, q = int(m.group(1)), int(m.group(2))
    return date(y, q * 3, (31, 30, 30, 31)[q - 1]).isoformat()


def _latest_workbook(dataset: dict):
    candidates = []
    for r in dataset.get("resources", []):
        title = r.get("title") or ""
        fmt = (r.get("format") or "").lower()
        if fmt != "xlsx" or "obs-hd-thd-deploiement" not in title.lower():
            continue
        period = _quarter(title)
        if period and (r.get("latest") or r.get("url")):
            candidates.append((period, r))
    if not candidates:
        raise RuntimeError("ARCEP deployment workbook not found")
    return max(candidates, key=lambda x: x[0])


def _national_ftth_total(wb):
    # Fail closed. ARCEP deployment workbooks are primarily geographic tables:
    # prefer an explicit national row, otherwise accept a single explicit
    # aggregate metric cell. Never sum geographic rows or operators here.
    label_rx = re.compile(r"(locaux|premises).*(raccordables|raccordable).*(ftth)|ftth.*(locaux|premises).*(raccordables|raccordable)", re.I)
    france_rx = re.compile(r"^(france|total france|ensemble france|france entière|total national|national)$", re.I)
    hits = []
    diagnostics = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for ri, row in enumerate(rows):
            cells = ["" if v is None else str(v).strip() for v in row]
            joined = " | ".join(cells)
            if not label_rx.search(joined):
                continue
            diagnostics.append({"sheet":ws.title,"row":ri+1,"context":joined[:500]})
            # Case 1: label and a single national aggregate value share a row.
            nums = [(ci, _safe_float(v)) for ci, v in enumerate(row)]
            nums = [(ci, v) for ci, v in nums if v is not None and v > 1_000_000]
            if len(nums) == 1 and any(france_rx.fullmatch(x) for x in cells):
                hits.append((ws.title, ri + 1, nums[0][0] + 1, nums[0][1], joined))
            # Case 2: explicit national row is near the metric header.
            for rj in range(max(0, ri - 8), min(len(rows), ri + 80)):
                rr = rows[rj]
                text = ["" if v is None else str(v).strip() for v in rr]
                if not any(france_rx.fullmatch(x) for x in text):
                    continue
                nums = [(ci, _safe_float(v)) for ci, v in enumerate(rr)]
                nums = [(ci, v) for ci, v in nums if v is not None and v > 1_000_000]
                if len(nums) == 1:
                    hits.append((ws.title, rj + 1, nums[0][0] + 1, nums[0][1], joined))
    unique = {(h[0], h[1], h[2], h[3]): h for h in hits}
    if len(unique) != 1:
        sample = [d for d in diagnostics if d["sheet"].lower() == "couverture"][:50]
        if not sample:\n            # Diagnostic-only: inspect actual Couverture rows regardless of wording.\n            for ws in wb.worksheets:\n                if ws.title.lower() != "couverture":\n                    continue\n                for ri, row in enumerate(ws.iter_rows(values_only=True)):\n                    cells = [str(v).strip() for v in row if v is not None and str(v).strip()]\n                    if not cells:\n                        continue\n                    sample.append({"sheet": ws.title, "row": ri + 1, "context": " | ".join(cells)[:700]})\n                    if len(sample) >= 50:\n                        break\n        if not sample:\n            sample = [{"sheet": ws.title, "max_row": ws.max_row, "max_column": ws.max_column} for ws in wb.worksheets]\n        raise RuntimeError(f"Expected one explicit national FTTH raccordable total, found {len(unique)}; diagnostic rows={sample}")
    return next(iter(unique.values()))


def load_arcep_deployment(ctx: PipelineContext) -> dict:
    source_id, run_id = start_run(ctx, "ARCEP_OBS", {"collector": "arcep_deployment_v5"})
    read = written = 0
    try:
        country = one(ctx.db, "countries", "iso3", "FRA")
        kpi = one(ctx.db, "kpis", "code", "FTTH_HOMES_PASSED")
        r = ctx.session.get(DATASET_API, timeout=45); r.raise_for_status()
        dataset = r.json()
        period, res = _latest_workbook(dataset)
        url = res.get("latest") or res.get("url")
        x = ctx.session.get(url, timeout=120); x.raise_for_status()
        wb = load_workbook(io.BytesIO(x.content), read_only=True, data_only=True)
        sheet, row, col, value, label = _national_ftth_total(wb)
        read = 1
        raw = {
            "ingestion_run_id": run_id, "source_id": source_id, "country_id": country["id"], "operator_id": None,
            "source_indicator": "Locaux raccordables FTTH - France", "period_date": period, "frequency": "quarterly",
            "value_numeric": value, "unit_raw": "locaux raccordables", "source_url": url, "retrieved_at": utcnow(),
            "payload": {"dataset_id": dataset.get("id"), "resource_id": res.get("id"), "resource_title": res.get("title"),
                        "sheet": sheet, "row": row, "column": col, "matched_context": label},
            "source_record_key": f"{res.get('id')}:{sheet}:{row}:{col}"
        }
        raw_id = _upsert_raw(ctx, raw)
        obs = {"kpi_id": kpi["id"], "country_id": country["id"], "operator_id": None, "period_date": period,
               "frequency": "quarterly", "value": value, "unit": "premises", "currency_code": None,
               "source_id": source_id, "raw_observation_id": raw_id, "definition_version": kpi["definition_version"],
               "quality_flag": "ok", "retrieved_at": utcnow(),
               "quality_notes": "ARCEP locaux raccordables FTTH; explicit national total, not summed across zones/operators."}
        _upsert_obs(ctx, obs); written = 1
        meta = {"collector": "arcep_deployment_v5", "period": period, "resource_id": res.get("id"), "source_url": url}
        finish_run(ctx, run_id, "success", read, written, metadata=meta)
        ctx.db.table("pipeline_state").upsert({"source_id": source_id, "last_success_at": utcnow(),
                                               "last_attempt_at": utcnow(), "cursor_state": meta}).execute()
        return {"rows_read": read, "rows_written": written, **meta}
    except Exception as exc:
        finish_run(ctx, run_id, "failed", read, written, str(exc)[:1000], {"collector": "arcep_deployment_v5"})
        raise
