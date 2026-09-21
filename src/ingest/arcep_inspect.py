from __future__ import annotations

import io

from openpyxl import load_workbook

from .arcep import DATASET_API
from .base import PipelineContext, finish_run, start_run


def inspect_arcep_workbooks(ctx: PipelineContext) -> dict:
    """Download ARCEP XLSX resources and persist workbook structure for parser development."""
    _, run_id = start_run(ctx, "ARCEP_OBS", {"collector": "arcep_xlsx_inspect_v1"})
    inspected = []
    try:
        response = ctx.session.get(DATASET_API, timeout=30)
        response.raise_for_status()
        dataset = response.json()
        resources = [r for r in dataset.get("resources", []) if (r.get("format") or "").lower() == "xlsx"]
        for resource in resources:
            url = resource.get("latest") or resource.get("url")
            response = ctx.session.get(url, timeout=60)
            response.raise_for_status()
            wb = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
            sheets = []
            for ws in wb.worksheets:
                preview = []
                for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True):
                    preview.append([None if v is None else str(v)[:180] for v in row[:12]])
                matches = []
                for ri, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    vals = ["" if v is None else str(v) for v in row[:4]]
                    joined = " | ".join(vals)
                    if any(term in joined.lower() for term in ("ftth", "fibre optique", "fiber")):
                        matches.append({"row": ri, "values": vals})
                sheets.append({"title": ws.title, "rows": ws.max_row, "cols": ws.max_column, "preview": preview, "ftth_matches": matches})
            inspected.append({"resource_id": resource.get("id"), "title": resource.get("title"),
                              "url": url, "bytes": len(response.content), "sheets": sheets})
        finish_run(ctx, run_id, "success", len(resources), 0, metadata={"collector": "arcep_xlsx_inspect_v1", "workbooks": inspected})
        return {"workbooks": len(inspected), "sheets": sum(len(x["sheets"]) for x in inspected)}
    except Exception as exc:
        finish_run(ctx, run_id, "failed", len(inspected), 0, str(exc)[:1000], {"collector": "arcep_xlsx_inspect_v1", "workbooks": inspected})
        raise
