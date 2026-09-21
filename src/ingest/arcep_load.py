from __future__ import annotations

import io
import re
from datetime import date

from openpyxl import load_workbook

from .arcep import DATASET_API
from .base import PipelineContext, finish_run, one, start_run, utcnow

ALIASES = {
    "CAPEX": ["Investment during the year (*)"],
    "MOBILE_SUBS": ["Total number of SIM cards (MtoM cards excluded)"],
    # ARCEP electronic-communications workbooks can expose FttH retail stock.
    # Keep aliases exact and fail closed: deployment/"locaux raccordables" data
    # belongs to the separate HD/THD deployment dataset and must not be mapped here.
    "FTTH_SUBS": [
        "Number of FttH subscriptions",
        "Number of subscriptions to FttH",
        "Number of FttH broadband subscriptions",
        "FttH broadband subscriptions",
        "FttH subscriptions",
        "of wich fiber",
        "of wich fiber to the home (FTTH)",
        "of which fiber to the home (FTTH)",
        "dont nombre d'abonnements en fibre de bout en bout (BLOM et BLOD)",
        "dont abonnements en fibre optique de bout en bout",
        "Nombre d'abonnements FttH",
        "Nombre d'abonnements en fibre optique de bout en bout (FttH)",
    ],
}


def _period(value, frequency):
    s = str(value).strip()
    if frequency == "annual" and re.fullmatch(r"20\d{2}|19\d{2}", s):
        return date(int(s), 12, 31).isoformat()
    m = re.fullmatch(r"Q([1-4])\s+(20\d{2})", s, re.I)
    if frequency == "quarterly" and m:
        q, y = int(m.group(1)), int(m.group(2))
        month = q * 3
        day = (31, 30, 30, 31)[q - 1]
        return date(y, month, day).isoformat()
    return None


def _scale(unit):
    u = (unit or "").lower()
    if "million" in u and ("unit" in u or "sim" in u or "subscription" in u or "abonnement" in u): return 1_000_000.0
    if "million" in u and ("€" in u or "eur" in u): return 1_000_000.0
    return 1.0


def _safe_float(v):
    if v is None or isinstance(v, bool): return None
    try: return float(v)
    except (TypeError, ValueError): return None


def _upsert_raw(ctx, payload):
    q = (ctx.db.table("raw_observations").select("id")
         .eq("source_id", payload["source_id"]).eq("country_id", payload["country_id"])
         .eq("source_indicator", payload["source_indicator"]).eq("period_date", payload["period_date"])
         .eq("frequency", payload["frequency"]).is_("operator_id", "null").limit(1).execute().data)
    if q:
        ctx.db.table("raw_observations").update(payload).eq("id", q[0]["id"]).execute()
        return q[0]["id"]
    return ctx.db.table("raw_observations").insert(payload).execute().data[0]["id"]


def _upsert_obs(ctx, payload):
    q = (ctx.db.table("observations").select("id")
         .eq("kpi_id", payload["kpi_id"]).eq("country_id", payload["country_id"])
         .eq("period_date", payload["period_date"]).eq("frequency", payload["frequency"])
         .eq("source_id", payload["source_id"]).is_("operator_id", "null").limit(1).execute().data)
    if q:
        ctx.db.table("observations").update(payload).eq("id", q[0]["id"]).execute()
    else:
        ctx.db.table("observations").insert(payload).execute()


def load_arcep(ctx: PipelineContext) -> dict:
    source_id, run_id = start_run(ctx, "ARCEP_OBS", {"collector": "arcep_load_v3"})
    country = one(ctx.db, "countries", "iso3", "FRA")
    kpis = {k: one(ctx.db, "kpis", "code", k) for k in ALIASES}
    read = written = 0
    matched = {}
    try:
        r = ctx.session.get(DATASET_API, timeout=30); r.raise_for_status(); dataset = r.json()
        resources = [x for x in dataset.get("resources", []) if (x.get("format") or "").lower() == "xlsx"]
        for res in resources:
            title = res.get("title") or ""
            if "DCOM" in title or "prix" in title.lower(): continue
            frequency = "quarterly" if "trimestriel" in title.lower() else "annual"
            url = res.get("latest") or res.get("url")
            x = ctx.session.get(url, timeout=60); x.raise_for_status()
            wb = load_workbook(io.BytesIO(x.content), read_only=True, data_only=True)
            for ws in wb.worksheets:
                if ws.title != "Open Data": continue
                rows = list(ws.iter_rows(values_only=True))
                header_idx = 1 if frequency == "quarterly" else 2
                periods = [_period(v, frequency) for v in rows[header_idx]]
                for ri, row in enumerate(rows):
                    label_en = str(row[0]).strip() if row and row[0] is not None else ""
                    label_fr = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
                    label = label_en or label_fr
                    for code, aliases in ALIASES.items():
                        if label_en not in aliases and label_fr not in aliases: continue
                        unit_en = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
                        unit_fr = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
                        unit = unit_en or unit_fr or kpis[code]["unit"]
                        matched[code] = matched.get(code, 0) + 1
                        for ci in range(4, min(len(row), len(periods))):
                            period = periods[ci]; val = _safe_float(row[ci])
                            if not period or val is None: continue
                            read += 1
                            numeric = val * _scale(unit)
                            raw = {
                                "ingestion_run_id": run_id, "source_id": source_id, "country_id": country["id"], "operator_id": None,
                                "source_indicator": label, "period_date": period, "frequency": frequency,
                                "value_text": str(row[ci]), "value_numeric": val, "unit_raw": unit,
                                "currency_raw": "EUR" if "€" in unit else None, "source_url": url,
                                "retrieved_at": utcnow(), "payload": {"sheet": ws.title, "row": ri + 1, "column": ci + 1,
                                "resource_id": res.get("id"), "resource_title": title},
                                "source_record_key": f"{res.get('id')}:{ws.title}:{ri+1}:{ci+1}"
                            }
                            raw_id = _upsert_raw(ctx, raw)
                            obs = {"kpi_id": kpis[code]["id"], "country_id": country["id"], "operator_id": None,
                                   "period_date": period, "frequency": frequency, "value": numeric,
                                   "unit": kpis[code]["unit"], "currency_code": "EUR" if code == "CAPEX" else None,
                                   "source_id": source_id, "raw_observation_id": raw_id,
                                   "definition_version": kpis[code]["definition_version"], "quality_flag": "ok",
                                   "retrieved_at": utcnow(), "quality_notes": f"ARCEP: {label}"}
                            _upsert_obs(ctx, obs); written += 1
        meta = {"collector":"arcep_load_v3","matched":matched}
        finish_run(ctx, run_id, "success", read, written, metadata=meta)
        ctx.db.table("pipeline_state").upsert({"source_id": source_id, "last_success_at": utcnow(), "last_attempt_at": utcnow(), "cursor_state": meta}).execute()
        return {"rows_read": read, "rows_written": written, "matched": matched}
    except Exception as exc:
        finish_run(ctx, run_id, "failed", read, written, str(exc)[:1000], {"collector":"arcep_load_v3","matched":matched})
        raise
