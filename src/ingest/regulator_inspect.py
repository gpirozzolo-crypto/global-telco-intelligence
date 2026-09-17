from __future__ import annotations

import io
import re
from urllib.parse import urljoin

from openpyxl import load_workbook

from .base import PipelineContext, finish_run, start_run

AGCOM_PAGE = "https://www.agcom.it/pubblicazioni/osservatori/osservatorio-sulle-comunicazioni-n-2-2026"
BNETZA_PAGE = "https://www.bundesnetzagentur.de/DE/Fachthemen/Telekommunikation/Marktdaten/artikel.html"
CNMC_PAGES = [
    "https://data.cnmc.es/telecomunicaciones-y-sector-audiovisual/conjuntos-de-datos/datos-mensuales/telecomunicaciones",
    "https://data.cnmc.es/telecomunicaciones-y-sector-audiovisual/datos-trimestrales/datos-generales/telecomunicaciones",
    "https://data.cnmc.es/telecomunicaciones-y-sector-audiovisual/datos-trimestrales/datos-de-mercados/telecomunicaciones-3",
]


def _workbook_preview(ctx: PipelineContext, url: str) -> dict:
    r = ctx.session.get(url, timeout=90, allow_redirects=True); r.raise_for_status()
    wb = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        preview = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15), values_only=True):
            preview.append([None if v is None else str(v)[:180] for v in row[:15]])
        sheets.append({"title": ws.title, "rows": ws.max_row, "cols": ws.max_column, "preview": preview})
    return {"url": r.url, "bytes": len(r.content), "sheets": sheets}


def inspect_agcom(ctx: PipelineContext) -> dict:
    _, run_id = start_run(ctx, "AGCOM_OBS", {"collector":"agcom_xlsx_inspect_v1"})
    try:
        r = ctx.session.get(AGCOM_PAGE, timeout=30); r.raise_for_status()
        links = re.findall(r'href=["\']([^"\']+\.xlsx[^"\']*)', r.text, re.I)
        links = list(dict.fromkeys(urljoin(r.url, x.replace("&amp;", "&")) for x in links))
        books = [_workbook_preview(ctx, u) for u in links[:3]]
        finish_run(ctx, run_id, "success", len(links), 0, metadata={"collector":"agcom_xlsx_inspect_v1","page":r.url,"workbooks":books})
        return {"workbooks":len(books),"links":len(links)}
    except Exception as exc:
        finish_run(ctx, run_id, "failed", 0, 0, str(exc)[:1000], {"collector":"agcom_xlsx_inspect_v1"}); raise


def inspect_bnetza(ctx: PipelineContext) -> dict:
    _, run_id = start_run(ctx, "BNetzA_TK", {"collector":"bnetza_xlsx_inspect_v1"})
    try:
        r = ctx.session.get(BNETZA_PAGE, timeout=30); r.raise_for_status()
        links = re.findall(r'href=["\']([^"\']+\.(?:xlsx|xlsm)[^"\']*)', r.text, re.I)
        links = list(dict.fromkeys(urljoin(r.url, x.replace("&amp;", "&")) for x in links))
        # Prefer the current annual-report data workbook.
        current = [u for u in links if "25_Daten_TK" in u] or links[:1]
        books = [_workbook_preview(ctx, u) for u in current[:2]]
        finish_run(ctx, run_id, "success", len(links), 0, metadata={"collector":"bnetza_xlsx_inspect_v1","page":r.url,"workbooks":books})
        return {"workbooks":len(books),"links":len(links)}
    except Exception as exc:
        finish_run(ctx, run_id, "failed", 0, 0, str(exc)[:1000], {"collector":"bnetza_xlsx_inspect_v1"}); raise


def inspect_cnmc(ctx: PipelineContext) -> dict:
    _, run_id = start_run(ctx, "CNMC_TELCO", {"collector":"cnmc_dataset_inspect_v1"})
    pages = []
    try:
        for url in CNMC_PAGES:
            r = ctx.session.get(url, timeout=30); r.raise_for_status()
            ids = list(dict.fromkeys(re.findall(r'[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}', r.text, re.I)))
            downloads = re.findall(r'href=["\']([^"\']+(?:csv|json|datastore_search)[^"\']*)', r.text, re.I)
            pages.append({"url":r.url,"bytes":len(r.content),"resource_ids":ids[:10],"download_links":[urljoin(r.url,x.replace("&amp;","&")) for x in downloads[:10]]})
        finish_run(ctx, run_id, "success", len(pages), 0, metadata={"collector":"cnmc_dataset_inspect_v1","pages":pages})
        return {"pages":len(pages),"resource_ids":sum(len(x["resource_ids"]) for x in pages)}
    except Exception as exc:
        finish_run(ctx, run_id, "failed", len(pages), 0, str(exc)[:1000], {"collector":"cnmc_dataset_inspect_v1","pages":pages}); raise
