from __future__ import annotations

import io
import re
from datetime import date, datetime

from openpyxl import load_workbook

from .base import PipelineContext, finish_run, one, start_run, utcnow
from .regulator_inspect import AGCOM_PAGE, BNETZA_PAGE


def _num(v):
    if v is None or isinstance(v, bool): return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace(" ", "").replace(",", ".")
    s = re.sub(r"[^0-9eE+.-]", "", s)
    try: return float(s)
    except ValueError: return None


def _year(v):
    m = re.search(r"(?:19|20)\d{2}", str(v or ""))
    return int(m.group(0)) if m else None


def _raw(ctx, p):
    q=(ctx.db.table("raw_observations").select("id").eq("source_id",p["source_id"]).eq("country_id",p["country_id"])
       .eq("source_indicator",p["source_indicator"]).eq("period_date",p["period_date"]).eq("frequency",p["frequency"])
       .is_("operator_id","null").limit(1).execute().data)
    if q:
        ctx.db.table("raw_observations").update(p).eq("id",q[0]["id"]).execute(); return q[0]["id"]
    return ctx.db.table("raw_observations").insert(p).execute().data[0]["id"]


def _obs(ctx, p):
    q=(ctx.db.table("observations").select("id").eq("kpi_id",p["kpi_id"]).eq("country_id",p["country_id"])
       .eq("period_date",p["period_date"]).eq("frequency",p["frequency"]).eq("source_id",p["source_id"])
       .is_("operator_id","null").limit(1).execute().data)
    if q: ctx.db.table("observations").update(p).eq("id",q[0]["id"]).execute()
    else: ctx.db.table("observations").insert(p).execute()


def _write(ctx, run_id, source_id, country, kpi, indicator, period, frequency, raw_value, raw_unit, value, source_url, payload, currency=None):
    raw_id=_raw(ctx,{"ingestion_run_id":run_id,"source_id":source_id,"country_id":country["id"],"operator_id":None,
        "source_indicator":indicator,"period_date":period,"frequency":frequency,"value_text":str(raw_value),"value_numeric":_num(raw_value),
        "unit_raw":raw_unit,"currency_raw":currency,"source_url":source_url,"retrieved_at":utcnow(),"payload":payload,
        "source_record_key":f"{payload.get('sheet')}:{indicator}:{period}"})
    _obs(ctx,{"kpi_id":kpi["id"],"country_id":country["id"],"operator_id":None,"period_date":period,"frequency":frequency,
        "value":value,"unit":kpi["unit"],"currency_code":currency,"source_id":source_id,"raw_observation_id":raw_id,
        "definition_version":kpi["definition_version"],"quality_flag":"ok","quality_notes":indicator,"retrieved_at":utcnow()})


def _xlsx_link(ctx, page, pattern):
    r=ctx.session.get(page,timeout=45); r.raise_for_status()
    links=re.findall(r'href=["\']([^"\']+\.(?:xlsx|xlsm)[^"\']*)',r.text,re.I)
    from urllib.parse import urljoin
    urls=[urljoin(r.url,x.replace("&amp;","&")) for x in links]
    selected=[u for u in urls if pattern in u]
    if not selected: raise RuntimeError(f"Workbook not found: {pattern}")
    return selected[0]


def load_agcom(ctx: PipelineContext) -> dict:
    source_id,run_id=start_run(ctx,"AGCOM_OBS",{"collector":"agcom_load_v1"}); read=written=0
    country=one(ctx.db,"countries","iso3","ITA")
    codes=["FIXED_BB_SUBS","FTTH_SUBS","MOBILE_SUBS","TELCO_REVENUE","FIXED_REVENUE","MOBILE_REVENUE","MOBILE_DATA_TRAFFIC"]
    k={c:one(ctx.db,"kpis","code",c) for c in codes}
    try:
        url=_xlsx_link(ctx,AGCOM_PAGE,"OPEN%20DATA%20Oss.%202-2026")
        x=ctx.session.get(url,timeout=90); x.raise_for_status(); wb=load_workbook(io.BytesIO(x.content),read_only=True,data_only=True)
        for sheet,row_idx,code,label in [("1.2",9,"FIXED_BB_SUBS","Totale - Total"),("1.7",8,"MOBILE_SUBS","Human (*)")]:
            ws=wb[sheet]; rows=list(ws.iter_rows(values_only=True)); dates=rows[2]
            for ci in range(1,min(9,len(dates))):
                d=dates[ci]; v=_num(rows[row_idx-1][ci])
                if not isinstance(d,(date,datetime)) or v is None: continue
                period=d.date().isoformat() if isinstance(d,datetime) else d.isoformat(); read+=1
                _write(ctx,run_id,source_id,country,k[code],label,period,"quarterly",v,"million",v*1_000_000,url,{"sheet":sheet,"row":row_idx,"column":ci+1}); written+=1
        ws=wb["Principali serie storiche"]; rows=list(ws.iter_rows(values_only=True)); headers=rows[2]
        ftth=next((r for r in rows if str(r[0]).strip()=="- FTTH"),None)
        if ftth:
            for ci in range(1,min(len(ftth),len(headers))):
                h=str(headers[ci] or ""); m=re.fullmatch(r"([1-4])T(\d{2})",h); v=_num(ftth[ci])
                if not m or v is None: continue
                q,y=int(m.group(1)),2000+int(m.group(2)); month=q*3; day=(31,30,30,31)[q-1]; period=date(y,month,day).isoformat(); read+=1
                _write(ctx,run_id,source_id,country,k["FTTH_SUBS"],"FTTH",period,"quarterly",v,"million",v*1_000_000,url,{"sheet":ws.title,"column":ci+1}); written+=1
        ws=wb["RA2026 1-2"]; rows=list(ws.iter_rows(values_only=True)); years=[int(v) if isinstance(v,(int,float)) else int(v) if str(v).isdigit() else None for v in rows[3]]
        rev={"Comunicazioni elettroniche":"TELCO_REVENUE","   - Rete fissa":"FIXED_REVENUE","   - Rete mobile":"MOBILE_REVENUE"}
        for r in rows:
            label=str(r[0] or "")
            if label not in rev: continue
            for ci,y in enumerate(years):
                v=_num(r[ci]) if ci<len(r) else None
                if not y or v is None: continue
                period=date(y,12,31).isoformat(); code=rev[label]; read+=1
                _write(ctx,run_id,source_id,country,k[code],label,period,"annual",v,"billion EUR",v*1_000_000_000,url,{"sheet":ws.title,"column":ci+1},"EUR"); written+=1
        meta={"collector":"agcom_load_v1","rows":written}; finish_run(ctx,run_id,"success",read,written,metadata=meta)
        ctx.db.table("pipeline_state").upsert({"source_id":source_id,"last_success_at":utcnow(),"last_attempt_at":utcnow(),"cursor_state":meta}).execute()
        return {"rows_read":read,"rows_written":written}
    except Exception as exc:
        finish_run(ctx,run_id,"failed",read,written,str(exc)[:1000],{"collector":"agcom_load_v1"}); raise


def load_bnetza(ctx: PipelineContext) -> dict:
    source_id,run_id=start_run(ctx,"BNetzA_TK",{"collector":"bnetza_load_v2"}); read=written=0
    country=one(ctx.db,"countries","iso3","DEU")
    codes=["TELCO_REVENUE","FIXED_REVENUE","MOBILE_REVENUE","CAPEX","FIXED_BB_SUBS","FTTH_SUBS","FTTH_HOMES_PASSED","FTTH_TAKEUP","MOBILE_SUBS","5G_SUBS","MOBILE_DATA_TRAFFIC"]
    k={c:one(ctx.db,"kpis","code",c) for c in codes}
    try:
        url=_xlsx_link(ctx,BNETZA_PAGE,"25_Daten_TK")
        x=ctx.session.get(url,timeout=90); x.raise_for_status(); wb=load_workbook(io.BytesIO(x.content),read_only=True,data_only=True)
        specs=[
          ("Außenumsatz TK-Markt","gesamt","TELCO_REVENUE",1_000_000_000,"billion EUR","EUR"),
          ("Investitionen TK-Markt","gesamt","CAPEX",1_000_000_000,"billion EUR","EUR"),
          ("Aktive BB-Anschlüsse Festnetz","gesamt","FIXED_BB_SUBS",1_000_000,"million",None),
          ("akt. FttH- und FttB-Anschlüsse","FttH","FTTH_SUBS",1_000_000,"million",None),
          ("Datenvolumen Mobil","Datenvolumen Mobilfunk insgesamt in Mio. GB¹⁾","MOBILE_DATA_TRAFFIC",1_000_000,"million GB",None),
        ]
        for sheet,label,code,scale,unit,currency in specs:
            rows=list(wb[sheet].iter_rows(values_only=True)); header=rows[1]
            target=next((r for r in rows if str(r[2] or "").strip()==label),None)
            if not target: continue
            for ci in range(3,min(len(target),len(header))):
                y=_year(header[ci]); v=_num(target[ci])
                if y is None or v is None: continue
                period=date(y,12,31).isoformat(); read+=1
                _write(ctx,run_id,source_id,country,k[code],label,period,"annual",v,unit,v*scale,url,{"sheet":sheet,"column":ci+1},currency); written+=1
        rows=list(wb["Außenumsatz Segmente"].iter_rows(values_only=True)); header=rows[1]
        for label,code in [("Außenumsatzerlöse über Festnetze","FIXED_REVENUE"),("Außenumsatzerlöse über Mobilfunknetze","MOBILE_REVENUE")]:
            target=next(r for r in rows if str(r[2] or "").strip()==label)
            for ci in (3,5,7):
                y=_year(header[ci]); v=_num(target[ci])
                if y is None or v is None: continue
                period=date(y,12,31).isoformat(); read+=1
                _write(ctx,run_id,source_id,country,k[code],label,period,"annual",v,"billion EUR",v*1_000_000_000,url,{"sheet":"Außenumsatz Segmente","column":ci+1},"EUR"); written+=1
        rows=list(wb["Glasfaser-Anschlüsse"].iter_rows(values_only=True)); header=rows[1]
        for label,code,scale,unit in [("Homes Passed","FTTH_HOMES_PASSED",1_000_000,"million"),("Take-up-Rate (FttH und FttB Activated bezogen auf Homes Passed)","FTTH_TAKEUP",100,"ratio")]:
            target=next(r for r in rows if str(r[2] or "").strip()==label)
            for ci in (3,4,5):
                y=_year(header[ci]); v=_num(target[ci])
                if y is None or v is None: continue
                period=date(y,12,31).isoformat(); read+=1
                _write(ctx,run_id,source_id,country,k[code],label,period,"annual",target[ci],unit,v*scale,url,{"sheet":"Glasfaser-Anschlüsse","column":ci+1}); written+=1
        rows=list(wb["SIM-Profile"].iter_rows(values_only=True)); years={4:2023,6:2024,8:2025}
        for label,code in [("insgesamt, ohne M2M","MOBILE_SUBS"),("davon 5G-Teilnehmer (NSA)","5G_SUBS")]:
            target=next(r for r in rows if str(r[2] or "").strip()==label)
            for ci,y in years.items():
                v=_num(target[ci]);
                if v is None: continue
                period=date(y,12,31).isoformat(); read+=1
                _write(ctx,run_id,source_id,country,k[code],label,period,"annual",v,"million",v*1_000_000,url,{"sheet":"SIM-Profile","column":ci+1}); written+=1
        meta={"collector":"bnetza_load_v2","rows":written}; finish_run(ctx,run_id,"success",read,written,metadata=meta)
        ctx.db.table("pipeline_state").upsert({"source_id":source_id,"last_success_at":utcnow(),"last_attempt_at":utcnow(),"cursor_state":meta}).execute()
        return {"rows_read":read,"rows_written":written}
    except Exception as exc:
        finish_run(ctx,run_id,"failed",read,written,str(exc)[:1000],{"collector":"bnetza_load_v2"}); raise
